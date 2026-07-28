"""Чтение «эталонного» реестра наименований из xlsx.

Файл-реестр (например «523-ПИР-24_СП_Наименование файлов.xlsx») содержит
для каждого документа: обозначение (шифр), раздел/подраздел/часть, статус
и ШАБЛОН имени файла (с подстрокой «ггммдд» вместо даты).

Из шаблона мы извлекаем:
    key            — короткий шифр-«хвост» для сопоставления (ИОС1.1, ПЗ, ПОС…)
    oboznachenie   — полное обозначение по стандарту (523-ПИР-24-ИОС1.1)
    middle         — фиксированная часть имени («Раздел ПД №5 Подраздел ПД №1 Часть №1»)
    doc_template   — шаблон имени документа
    iul_template   — шаблон имени ИУЛа
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional

try:
    import openpyxl
except ImportError:  # подсказка при отсутствии зависимости
    openpyxl = None


# Заголовки колонок, которые ищем в строке-шапке (регистр/пробелы не важны)
COL_ALIASES = {
    "stage": ["стадия", "п/и", "пд/ии", "стадия (п/и)"],
    "razdel": ["раздел"],
    "podrazdel": ["подраздел"],
    "chast": ["часть"],
    "kniga": ["книга", "книги"],
    "type": ["тип", "тип документа", "тип сметы", "вид документа"],
    "izy_part": ["часть изысканий", "часть изыск", "часть ии"],
    "nomer": ["номер тома", "номер", "том"],
    "short": ["шифр раздела", "шифр"],
    "oboznachenie": ["обозначение"],
    "name": ["наименование"],
    "org": ["исполнитель", "организация"],
    "status": ["статус"],
    "doc_file": ["наименование файла документации", "файл документации"],
    "iul_file": ["наименование файла иула", "файл иула", "файл иул"],
}

# Типы сметной документации (в колонке «Тип») — попадают в раздел 12_СМ
SMETA_TYPES = ["ЛСР", "ОСР", "ВОР", "СВОР", "ССР", "ССРСС", "СР", "КАЦ", "ПЗ", "ПИР"]

DATE_TOKEN = "ггммдд"

# Форма шифра объекта задаётся ШАБЛОНОМ (см. configure_shifr) — по умолчанию
# «число-буквы-число» (523-ПИР-24). Это позволяет работать с любым форматом
# (728-ПСД-25 и т.п.) без правки кода — достаточно поменять шаблон в настройках.
DEFAULT_SHIFR_PATTERN = r"\d+[-_][A-Za-zА-Яа-яЁё]+[-_]\d+"
SHIFR_CORE = re.compile(DEFAULT_SHIFR_PATTERN, re.IGNORECASE)        # сам шифр
RE_PROJ_PREFIX = re.compile(rf"^\s*(?:{DEFAULT_SHIFR_PATTERN})\s*[-_]\s*", re.IGNORECASE)


def configure_shifr(pattern: str | None) -> None:
    """Задать форму шифра объекта (регулярное выражение). Пересобирает шаблоны,
    которыми распознаётся шифр и отрезается префикс документа."""
    global SHIFR_CORE, RE_PROJ_PREFIX
    pat = (pattern or "").strip() or DEFAULT_SHIFR_PATTERN
    try:
        SHIFR_CORE = re.compile(pat, re.IGNORECASE)
        RE_PROJ_PREFIX = re.compile(rf"^\s*(?:{pat})\s*[-_]\s*", re.IGNORECASE)
    except re.error:                       # некорректный шаблon — откатываемся
        SHIFR_CORE = re.compile(DEFAULT_SHIFR_PATTERN, re.IGNORECASE)
        RE_PROJ_PREFIX = re.compile(rf"^\s*(?:{DEFAULT_SHIFR_PATTERN})\s*[-_]\s*", re.IGNORECASE)


def strip_proj_prefix(s: str) -> str:
    """523-ПИР-24-ИОС1.1 -> ИОС1.1 ; 560-ПИР-24-СМ4-ВОР1 -> СМ4-ВОР1.
    Форма шифра задаётся configure_shifr (по умолчанию число-буквы-число)."""
    return RE_PROJ_PREFIX.sub("", str(s or "").strip())


@dataclass
class RegEntry:
    oboznachenie: str            # 523-ПИР-24-ИОС1.1
    key: str                     # ИОС1.1  (для сопоставления с файлами)
    razdel: str = ""
    podrazdel: str = ""
    chast: str = ""
    short: str = ""              # ИОС
    tom: str = ""                # «Номер тома» (для титульных листов)
    name: str = ""
    status: str = ""
    doc_template: str = ""       # ..._ггммдд.pdf
    iul_template: str = ""       # ..._ггммдд_УЛ.pdf
    extensions: List[str] = field(default_factory=list)  # допускаемые типы в !PUBLISHED
    signers: List[str] = field(default_factory=list)  # обязательные подписанты ЭЦП
    signature_level: str = "warn"  # none / warn / error
    area: str = "ПД"             # "ПД" или "ИИ" — область документации
    section: str = ""            # явная папка-раздел (для ИИ: 01_ИГДИ); пусто = вычислить
    # имена подпапок каталога версии (по умолчанию берутся из настроек проекта).
    # split=True — том делится на !EDIT/!PUBLISHED (программа их требует);
    # split=False — версия хранится одной папкой (сметы: связи в xlsx).
    split: bool = True
    edit_dir: str = ""           # имя папки редактируемых форматов (пусто = из настроек)
    pub_dir: str = ""            # имя папки публикуемых файлов (пусто = из настроек)

    def middle(self) -> str:
        """Средняя часть имени между обозначением и датой.
        Учитывает, что само обозначение может содержать «_» (ЛСР/ОСР)."""
        t = self.doc_template
        if not t:
            return ""
        body = t.rsplit(".", 1)[0]                      # без расширения
        prefix = self.oboznachenie + "_"
        if body.startswith(prefix):
            body = body[len(prefix):]
        body = re.sub(r"_?" + re.escape(DATE_TOKEN) + r"$", "", body,
                      flags=re.IGNORECASE)              # убрать дату
        return body.strip("_")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _tail(oboznachenie: str) -> str:
    """523-ПИР-24-ИОС1.1 -> ИОС1.1 (хвост после последнего дефиса)."""
    o = str(oboznachenie or "").strip()
    return o.split("-")[-1] if "-" in o else o


def read_object_name(path: str, sheet: Optional[str] = None) -> str:
    """Наименование объекта из верхних строк реестра (для титульных листов).
    Ищет ячейку «Наименование объекта» и берёт значение справа."""
    if openpyxl is None:
        return ""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001  — openpyxl кидает разные типы; поле декоративное
        return ""
    _ws, rows = _pick_sheet(wb, sheet)
    for r in rows[:8]:
        for j, c in enumerate(r):
            if c and "наименование объект" in str(c).strip().lower():
                for k in range(j + 1, len(r)):
                    if r[k] and str(r[k]).strip():
                        return str(r[k]).strip().strip("«»\"")
    return ""


def _pick_sheet(wb, sheet: Optional[str]):
    """Берёт указанный лист; если его нет или в нём не находится шапка с «Раздел» —
    автоматически ищет подходящий лист (ПИР/СП/первый с колонкой «Раздел»)."""
    def _rows(ws):
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def _has_header(rows):
        for r in rows[:8]:
            if any("раздел" == str(c).strip().lower() for c in r if c):
                return True
        return False

    if sheet and sheet in wb.sheetnames:
        rows = _rows(wb[sheet])
        if _has_header(rows):
            return wb[sheet], rows
    for name in wb.sheetnames:                       # ищем лист с шапкой
        rows = _rows(wb[name])
        if _has_header(rows):
            return wb[name], rows
    ws = wb[wb.sheetnames[0]]
    return ws, _rows(ws)


def _strip_date(name: str) -> str:
    """Убирает из имени файла версию-дату «_ггммдд» или «_260601» (перед
    расширением или в конце)."""
    if not name:
        return name
    s = re.sub(r"_(?:" + re.escape(DATE_TOKEN) + r"|\d{6})(?=\.[^.]+$|$)", "",
               str(name).strip(), flags=re.IGNORECASE)
    s = re.sub(r"_+(\.[^.]+)$", r"\1", s)   # убрать висящий «_» перед расширением
    s = re.sub(r"_+$", "", s)               # …или в самом конце
    return s.strip()


def parse_row(r, colmap, proj, smeta_sub=("ВОР", "ЛСР", "ОСР"), smeta_razdel="12"):
    """Разбирает одну строку реестра в RegEntry или None (если это не документ:
    служебная строка-заголовок, изыскание без раздела и т.п.). Используется и
    при загрузке, и при конвертации реестров в стандартную форму."""
    def cell(field):
        j = colmap.get(field)
        return "" if (j is None or j >= len(r) or r[j] is None) else str(r[j]).strip()

    # строки изысканий (Стадия = И) в состав ПД не идут — изыскания берутся с диска
    if cell("stage").strip().upper().startswith("И"):
        return None
    razdel = _numstr(cell("razdel"))
    podr = _numstr(cell("podrazdel"))
    chast = _numstr(cell("chast"))
    kniga = _numstr(cell("kniga"))
    short = cell("short")
    raw = cell("oboznachenie")
    short_u = short.upper()
    sub_codes = {s.upper() for s in smeta_sub}
    smeta_set = {t.upper() for t in SMETA_TYPES}
    smtype = cell("type").upper()
    if smtype not in smeta_set and short_u in {"ЛСР", "ОСР", "ВОР", "СВОР",
                                               "ССРСС", "КАЦ"}:
        smtype = short_u

    part = chast or "3"
    if smtype in ("ЛСР", "ОСР"):
        code = raw or cell("nomer")
        oboz = f"{proj}-СМ{part}_{code}_{smtype}" if (proj and code) else raw
    elif smtype:
        oboz = raw or (f"{proj}-СМ{part}_{smtype}" if proj else smtype)
    elif raw and re.fullmatch(r"\d+-\d+(?:-\d+)?", raw):
        typ = "ЛСР" if raw.count("-") == 2 else "ОСР"
        oboz = f"{proj}-СМ{part}_{raw}_{typ}" if proj else raw
    else:
        code = cell("nomer")
        if not raw and code and re.fullmatch(r"\d+-\d+(?:-\d+)?", code):
            typ = "ЛСР" if code.count("-") == 2 else "ОСР"
            oboz = f"{proj}-СМ{part}_{code}_{typ}" if proj else code
        else:
            oboz = raw or _gen_oboz(proj, short, podr, chast, kniga)

    is_sub = bool(smtype) or short_u in sub_codes or _has_sub(oboz, sub_codes)
    if not razdel and is_sub:
        razdel = str(smeta_razdel)
    if not oboz or not razdel:
        return None
    key = strip_proj_prefix(oboz)
    if not key:
        return None

    doc_file = cell("doc_file")
    if kniga:
        doc_template = f"{oboz}_{_build_middle(razdel, podr, chast, kniga)}.pdf"
    elif doc_file:
        doc_template = _strip_date(doc_file)
    else:
        doc_template = f"{oboz}_{_build_middle(razdel, podr, chast, kniga)}.pdf"
    iul_file = cell("iul_file")
    iul_template = _strip_date(iul_file) if iul_file else ""
    section_short = "СМ" if is_sub else (short or _short_from_oboz(oboz))
    tom = cell("nomer")
    if tom.upper() in smeta_set:
        tom = ""
    return RegEntry(
        oboznachenie=oboz, key=key, razdel=razdel, podrazdel=podr, chast=chast,
        short=section_short, tom=tom,
        name=re.sub(r"\s+", " ", cell("name")).strip(), status=cell("status"),
        doc_template=doc_template, iul_template=iul_template)


def load_registry(path: str, sheet: Optional[str] = None,
                  smeta_sub=("ВОР", "ЛСР", "ОСР"), smeta_razdel="12") -> List[RegEntry]:
    entries, _ = load_registry_diag(path, sheet, smeta_sub, smeta_razdel)
    return entries


def load_registry_diag(path: str, sheet: Optional[str] = None,
                       smeta_sub=("ВОР", "ЛСР", "ОСР"), smeta_razdel="12"):
    """Возвращает (список_документов, строка_диагностики).
    Бросает понятную ошибку, если реестр пуст или нечитаем."""
    if openpyxl is None:
        raise RuntimeError("Не установлен модуль openpyxl (нужен для чтения xlsx).")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(
            f"Не удалось открыть файл реестра:\n{path}\n\n{e}\n\n"
            "Проверьте, что выбран правильный .xlsx и он полностью "
            "синхронизирован (не значок-заглушка облака).")
    sheets = wb.sheetnames
    ws, rows = _pick_sheet(wb, sheet)

    header_idx, colmap = _find_header(rows)
    data_start = (header_idx + 1) if header_idx is not None else 0
    proj = _project_shifr(rows, colmap, data_start)

    entries: List[RegEntry] = []
    seen = set()
    for r in rows[data_start:]:
        e = parse_row(r, colmap, proj, smeta_sub, smeta_razdel)
        if e is None or e.key.lower() in seen:
            continue
        seen.add(e.key.lower())
        entries.append(e)

    diag = (f"лист «{ws.title}», строк {len(rows)}, шифр объекта «{proj or '?'}», "
            f"колонки: раздел={_col_letter(colmap.get('razdel'))}, "
            f"шифр={_col_letter(colmap.get('short'))}, "
            f"обозначение={_col_letter(colmap.get('oboznachenie'))}; "
            f"документов {len(entries)}")
    if not entries:
        raise RuntimeError(
            "Реестр прочитан, но не удалось собрать ни одного документа.\n\n"
            f"Файл: {path}\nЛисты: {', '.join(sheets)}\nВыбранный лист: {ws.title}\n"
            f"Строк: {len(rows)}\n"
            f"Колонка «Раздел»: {_col_letter(colmap.get('razdel'))}\n"
            f"Колонка «ШИФР»: {_col_letter(colmap.get('short'))}\n"
            f"Колонка «Обозначение»: {_col_letter(colmap.get('oboznachenie'))}\n\n"
            "Нужны колонки: Раздел, Подраздел, Часть, ШИФР, Обозначение. "
            "Проверьте имя листа в Настройках (обычно «ПИР»).")
    return entries, diag


def _numstr(s) -> str:
    """«1.0» → «1», «2.1» оставляем, пустое → «»."""
    s = str(s if s is not None else "").strip()
    if not s:
        return ""
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return s


def _has_sub(oboz: str, sub_codes) -> bool:
    tail = strip_proj_prefix(oboz).upper()
    return any(s in tail for s in sub_codes)


def _build_middle(razdel: str, podr: str, chast: str, kniga: str = "") -> str:
    """«Раздел ПД №5 Подраздел ПД №1 Часть №1 Книга №2»."""
    m = f"Раздел ПД №{razdel}"
    if podr:
        m += f" Подраздел ПД №{podr}"
    if chast:
        m += f" Часть №{chast}"
    if kniga:
        m += f" Книга №{kniga}"
    return m


def _gen_oboz(proj: str, short: str, podr: str, chast: str, kniga: str = "") -> str:
    """Обозначение из шифра объекта и номеров: 523-ПИР-24-ИОС1.1."""
    if not (proj and short):
        return ""
    tail = short
    if podr and chast:
        tail = f"{short}{podr}.{chast}"
    elif chast:
        tail = f"{short}{chast}"
    elif podr:
        tail = f"{short}{podr}"
    if kniga:
        tail = f"{tail}.{kniga}"
    return f"{proj}-{tail}"


def _project_shifr(rows, colmap, data_start: int) -> str:
    """Ищем шифр объекта вида 523-ПИР-24 в верхних строках или по обозначению."""
    pat = re.compile(r"\d+-[A-Za-zА-Яа-яЁё]+-\d+")
    for r in rows[:8]:
        for c in r:
            s = str(c or "").strip()
            if re.fullmatch(pat, s):
                return s
    j = colmap.get("oboznachenie")
    if j is not None:
        for r in rows[data_start:]:
            if j < len(r) and r[j]:
                m = re.match(r"(\d+-[^-]+-\d+)-", str(r[j]).strip())
                if m:
                    return m.group(1)
    return ""


def _col_letter(idx) -> str:
    if idx is None:
        return "—"
    s = ""
    n = idx + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _short_from_oboz(oboz: str) -> str:
    """ПЗ из 523-ПИР-24-ПЗ; ИОС из 523-ПИР-24-ИОС1.1."""
    tail = strip_proj_prefix(oboz).split("-")[0]
    m = re.match(r"[A-Za-zА-Яа-яЁё]+", tail)
    return m.group(0) if m else tail


def _clean_template(t: str) -> str:
    t = re.sub(r"\s+", " ", str(t or "")).strip()
    return t if DATE_TOKEN in t.lower() else ""


def _detect_template_columns(rows, data_start: int):
    """Находит колонки шаблонов по содержимому: ячейки с «ггммдд».
    Колонка ИУЛа — где в шаблонах есть _УЛ/ИУЛ, документа — без него."""
    doc_score: Dict[int, int] = {}
    iul_score: Dict[int, int] = {}
    for r in rows[data_start:]:
        for j, c in enumerate(r):
            s = str(c or "")
            if DATE_TOKEN in s.lower() and s.lower().endswith((".pdf", ".sig", ".docx", ".xlsx", "ггммдд", "ггммдд_ул")):
                if re.search(r"_ул|иул", s, re.IGNORECASE):
                    iul_score[j] = iul_score.get(j, 0) + 1
                else:
                    doc_score[j] = doc_score.get(j, 0) + 1
    # запасной вариант: вообще любые ячейки с ггммдд
    if not doc_score and not iul_score:
        for r in rows[data_start:]:
            for j, c in enumerate(r):
                s = str(c or "")
                if DATE_TOKEN in s.lower():
                    (iul_score if re.search(r"_ул|иул", s, re.I) else doc_score)[j] = 1
    doc_col = max(doc_score, key=doc_score.get) if doc_score else None
    iul_col = max(iul_score, key=iul_score.get) if iul_score else None
    return doc_col, iul_col


def _find_header(rows) -> tuple[Optional[int], Dict[str, int]]:
    """Ищем строку-шапку в первых 30 строках; колонки собираем из «полосы»
    заголовков (заголовки могут занимать 1–3 строки)."""
    best = None
    for i, r in enumerate(rows[:30]):
        norm = [_norm(c) for c in r]
        if any("обознач" in c for c in norm) or any("наименование файла" in c for c in norm):
            best = i if best is None else best
    if best is None:
        return None, {}
    colmap: Dict[str, int] = {}
    band = rows[best:best + 3]
    for r in band:
        norm = [_norm(c) for c in r]
        for field, aliases in COL_ALIASES.items():
            if field in colmap:
                continue
            for j, c in enumerate(norm):
                if any(c == a or c.startswith(a) for a in aliases):
                    colmap[field] = j
                    break
    return best, colmap


def index_by_key(entries: List[RegEntry]) -> Dict[str, RegEntry]:
    return {e.key.lower(): e for e in entries}


def section_code(razdel, short: str, razdel_codes=None) -> str:
    """Имя папки раздела: 01_ПЗ, 05_ИОС, 12_СМ, 13_ИД.
    Для разделов из razdel_codes имя берётся оттуда (раздел 13 → ИД),
    иначе — из шифра документа."""
    razdel_codes = razdel_codes or {}
    rs = str(razdel or "").strip()
    try:
        n = int(float(rs))
        num, key = f"{n:02d}", str(n)
    except (ValueError, TypeError):
        return short or rs or "Прочее"
    code = razdel_codes.get(key) or razdel_codes.get(rs) or short or "Раздел"
    return f"{num}_{code}"
