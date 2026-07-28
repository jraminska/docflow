"""Формирование комплекта для передачи заказчику из !LATEST."""
from __future__ import annotations

import os
import shutil
import zlib
from datetime import datetime
from typing import Iterable, List

from .config import AppConfig
from . import planner

REGISTRY_NAME = "Реестр_переданной_документации.xlsx"
SERVICE_FILES = {".docflow_crc.json", ".docflow_latest.json", REGISTRY_NAME.lower()}


def _crc32(path: str, chunk: int = 1 << 20) -> str:
    crc = 0
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(chunk), b""):
            crc = zlib.crc32(block, crc)
    return f"{crc & 0xFFFFFFFF:08X}"


def _scan_roots(source_root: str, selected_folders) -> List[str]:
    if selected_folders is None:
        return [source_root]
    roots = []
    for item in selected_folders:
        path = item if os.path.isabs(item) else os.path.join(source_root, item)
        path = os.path.normpath(path)
        if os.path.isdir(path):
            roots.append(path)
    return roots


def _entry_index(entries) -> dict:
    result = {}
    for entry in entries or []:
        for value in (entry.oboznachenie, entry.key):
            if value:
                result[value.casefold().replace(" ", "")] = entry
    return result


def _entry_order(entries) -> dict:
    return {id(entry): index for index, entry in enumerate(entries or [])}


def _area_order(area: str) -> int:
    value = (area or "").upper()
    if value == "ПД":
        return 0
    if value == "ИИ":
        return 1
    return 2


def _area_title(area: str) -> str:
    if (area or "").upper() == "ПД":
        return "Проектная документация"
    if (area or "").upper() == "ИИ":
        return "Отчетная техническая документация"
    return area or "Документация"


def _file_context(rel: str, entry_lookup: dict) -> tuple:
    parts = rel.replace("\\", "/").split("/")
    entry = None
    version = ""
    designation = ""
    for part in reversed(parts[:-1]):
        prefix, date = planner.cat_match(part)
        if prefix:
            version = date or ""
            designation = prefix
            entry = entry_lookup.get(prefix.casefold().replace(" ", ""))
            break
    section = parts[1] if len(parts) > 2 else parts[0]
    if entry is not None:
        section = entry.section or section
        designation = entry.oboznachenie or designation
    volume_name = (
        getattr(entry, "name", "") if entry is not None else ""
    ) or designation or (parts[-2] if len(parts) > 1 else "")
    volume_number = getattr(entry, "tom", "") if entry is not None else ""
    razdel = getattr(entry, "razdel", "") if entry is not None else ""
    area = getattr(entry, "area", "") if entry is not None else ""
    return (
        entry, section, designation, volume_name, version,
        volume_number, razdel, area)


def _section_title(entry, section: str) -> str:
    name = (getattr(entry, "name", "") if entry is not None else "").strip()
    if name:
        return name.splitlines()[0]
    return section


def _signer_names(entry) -> List[str]:
    names = []
    for signer in getattr(entry, "signers", []) if entry is not None else []:
        if isinstance(signer, dict):
            signer = signer.get("name") or signer.get("fio") or ""
        signer = str(signer).strip()
        if signer:
            names.append(signer)
    return names


def _signature_signer(filename: str, entry) -> str:
    folded = filename.casefold().replace("ё", "е")
    for signer in _signer_names(entry):
        surname = signer.split()[0].casefold().replace("ё", "е")
        if surname and surname in folded:
            return signer
    return ""


def build_package(cfg: AppConfig, source_root: str, target_dir: str,
                  selected_folders=None, progress=None, entries=None) -> dict:
    """Скопировать выбранные ветки с сохранением структуры и создать реестр."""
    if not source_root or not os.path.isdir(source_root):
        raise FileNotFoundError("Папка !LATEST не найдена.")
    if not target_dir:
        raise ValueError("Не указана папка для комплекта.")
    os.makedirs(target_dir, exist_ok=True)
    rows = []
    copied = 0
    entry_lookup = _entry_index(entries)
    entry_positions = _entry_order(entries)
    for scan_root in _scan_roots(source_root, selected_folders):
        for dirpath, dirnames, filenames in os.walk(scan_root):
            dirnames[:] = [
                d for d in dirnames
                if d != cfg.archive_name and not planner._ignored(d, cfg)]
            for filename in sorted(filenames, key=str.casefold):
                if (filename.casefold() in SERVICE_FILES
                        or planner._ignored(filename, cfg)):
                    continue
                src = os.path.join(dirpath, filename)
                rel = os.path.relpath(src, source_root)
                dst = os.path.join(target_dir, rel)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                shutil.copy2(src, dst)
                size = os.path.getsize(dst)
                crc = _crc32(dst)
                (
                    entry, section, designation, volume_name, version,
                    volume_number, razdel, area,
                ) = _file_context(rel, entry_lookup)
                is_signature = filename.casefold().endswith(".sig")
                rows.append({
                    "section": section,
                    "section_title": _section_title(entry, section),
                    "group_key": f"{area}:{razdel or section}",
                    "area": area,
                    "entry_order": entry_positions.get(id(entry), 10 ** 9),
                    "designation": designation,
                    "volume_name": volume_name,
                    "volume_number": volume_number,
                    "version": version,
                    "relative": rel,
                    "folder": os.path.dirname(rel),
                    "name": filename,
                    "extension": os.path.splitext(filename)[1].lower().lstrip("."),
                    "size": size,
                    "crc32": crc,
                    "kind": "Подпись" if is_signature else "Документ",
                    "signer": _signature_signer(filename, entry) if is_signature else "",
                    "required_signers": ", ".join(_signer_names(entry)),
                })
                copied += 1
                if progress:
                    progress(copied, rel)
    registry = os.path.join(target_dir, REGISTRY_NAME)
    rows.sort(key=lambda row: (
        _area_order(row["area"]), row["entry_order"],
        row["designation"].casefold(),
        row["version"], row["name"].casefold().replace(".sig", ""),
        row["kind"] == "Подпись"))
    _write_registry(
        registry, rows, cfg.object_name, planner._project_shifr(entries or []))
    return {
        "target": target_dir,
        "registry": registry,
        "files": copied,
        "bytes": sum(row["size"] for row in rows),
        "rows": rows,
    }


def _write_registry(
        path: str, rows: Iterable[dict], project_name: str, project_code: str
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError(
            "Не установлен модуль openpyxl — невозможно создать Excel-реестр."
        ) from exc
    rows = list(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Состав"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.merge_cells("B1:G1")
    ws["A1"] = "Название проекта:"
    ws["B1"] = project_name or "—"
    ws.merge_cells("B2:G2")
    ws["A2"] = "Шифр объекта:"
    ws["B2"] = project_code or "—"
    ws.merge_cells("B3:G3")
    ws["A3"] = "Дата передачи:"
    ws["B3"] = datetime.now().strftime("%d.%m.%Y")
    for cell in ("A1", "A2", "A3"):
        ws[cell].font = Font(name="Arial", size=11, bold=True)
    for cell in ("B1", "B2", "B3"):
        ws[cell].font = Font(name="Arial", size=11, bold=True, color="1F4E78")
        ws[cell].alignment = Alignment(horizontal="left", vertical="center")
    headers = [
        "Номер тома", "Обозначение", "Наименование", "Версия",
        "Название файла", "CRC32", "Путь к файлу"]
    ws.append([])
    ws.append([])
    ws.append(headers)
    header_row = 6
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9E2F3")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[header_row].height = 42

    previous_area = None
    previous_group = None
    for row in rows:
        if row["area"] != previous_area:
            ws.append(["", "", _area_title(row["area"]), "", "", "", ""])
            section_row = ws.max_row
            for cell in ws[section_row]:
                cell.fill = section_fill
                cell.font = Font(
                    name="Arial", size=10, bold=True,
                    color="44546A")
                cell.border = border
            previous_area = row["area"]
            previous_group = None
        if (row["area"] or "").upper() == "ПД" and row["group_key"] != previous_group:
            ws.append(["", "", row["section_title"], "", "", "", ""])
            section_row = ws.max_row
            for cell in ws[section_row]:
                cell.fill = section_fill
                cell.font = Font(
                    name="Arial", size=10, bold=True,
                    color="44546A")
                cell.border = border
            previous_group = row["group_key"]
        ws.append([
            row["volume_number"], row["designation"], row["volume_name"],
            row["version"], row["name"], row["crc32"], row["folder"]])
        excel_row = ws.max_row
        link = ws.cell(excel_row, 7)
        link.hyperlink = row["folder"] or "."
        link.style = "Hyperlink"
        for cell in ws[excel_row]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        ws.cell(excel_row, 1).alignment = Alignment(
            horizontal="right", vertical="center")
        ws.cell(excel_row, 4).alignment = Alignment(
            horizontal="right", vertical="center")
        ws.cell(excel_row, 6).number_format = "@"
    widths = {
        "A": 22, "B": 25, "C": 55, "D": 13,
        "E": 55, "F": 14, "G": 65}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f"A{header_row}:G{max(header_row, ws.max_row)}"
    ws.print_title_rows = f"1:{header_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_orientation = "landscape"
    wb.save(path)
