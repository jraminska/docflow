from __future__ import annotations

import os
import sys
import zlib

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from docflow.config import AppConfig  # noqa: E402
from docflow import customer_package  # noqa: E402
from docflow.registry import RegEntry  # noqa: E402


def test_customer_package_preserves_structure_and_creates_crc_registry(tmp_path):
    source = tmp_path / "!LATEST"
    pz = source / "02_ПД" / "01_ПЗ" / "523-ПИР-24-ПЗ" / "523-ПИР-24-ПЗ_260728"
    kr = source / "02_ПД" / "04_КР" / "523-ПИР-24-КР1" / "523-ПИР-24-КР1_260728"
    archive = pz / "!ARCHIVE"
    pz.mkdir(parents=True)
    kr.mkdir(parents=True)
    archive.mkdir()
    content = b"<xml/>"
    (pz / "note.xml").write_bytes(content)
    (pz / "note.xml_Иванов.sig").write_bytes(b"sig")
    (kr / "drawing.pdf").write_bytes(b"pdf")
    (archive / "old.pdf").write_bytes(b"old")
    target = tmp_path / "delivery" / "customer" / "new-folder"
    cfg = AppConfig(
        project_root=str(tmp_path),
        latest_dir="!LATEST",
        object_name="Тестовый объект")
    entries = [RegEntry(
        oboznachenie="523-ПИР-24-ПЗ",
        key="ПЗ",
        section="01_ПЗ",
        name="Раздел 1. Пояснительная записка",
        tom="1",
        signers=["Иванов Иван Иванович"],
    )]

    result = customer_package.build_package(
        cfg, str(source), str(target),
        selected_folders=[os.path.join("02_ПД", "01_ПЗ", "523-ПИР-24-ПЗ")],
        entries=entries)

    assert result["files"] == 2
    copied = target / "02_ПД" / "01_ПЗ" / "523-ПИР-24-ПЗ" / "523-ПИР-24-ПЗ_260728"
    assert (copied / "note.xml").read_bytes() == content
    assert not (target / "02_ПД" / "04_КР").exists()
    assert not (copied / "!ARCHIVE").exists()

    wb = load_workbook(result["registry"])
    ws = wb["Состав"]
    assert ws.freeze_panes == "A7"
    assert ws["A1"].value == "Название проекта:"
    assert ws["B1"].value == "Тестовый объект"
    assert ws["B2"].value == "523-ПИР-24"
    assert ws["A3"].value == "Дата передачи:"
    assert isinstance(ws["B3"].value, str)
    assert ws["D6"].value == "Версия"
    assert ws["E6"].value == "Название файла"
    assert ws["F6"].value == "CRC32"
    assert ws["G6"].value == "Путь к файлу"
    names = [ws.cell(row, 5).value for row in range(7, ws.max_row + 1)
             if ws.cell(row, 5).value]
    assert names == ["note.xml", "note.xml_Иванов.sig"]
    crc_row = next(
        row for row in range(7, ws.max_row + 1)
        if ws.cell(row, 5).value == "note.xml")
    assert ws.cell(crc_row, 1).value == "1"
    assert ws.cell(crc_row, 2).value == "523-ПИР-24-ПЗ"
    assert ws.cell(crc_row, 3).value == "Раздел 1. Пояснительная записка"
    assert ws.cell(crc_row, 4).value == "260728"
    assert ws.cell(crc_row, 6).value == f"{zlib.crc32(content) & 0xFFFFFFFF:08X}"
    assert ws.cell(crc_row, 7).hyperlink is not None
    assert ws.cell(crc_row, 7).hyperlink.target.endswith(
        os.path.join("523-ПИР-24-ПЗ", "523-ПИР-24-ПЗ_260728"))
    wb.close()


def test_registry_puts_pd_before_surveys_and_keeps_volume_order(tmp_path):
    source = tmp_path / "!LATEST"
    files = [
        ("01_ИИ", "02_ИГИ", "523-ПИР-24-ИГИ", "survey.pdf"),
        ("02_ПД", "02_СПОЗУ", "523-ПИР-24-ПЗУ1", "volume-2.pdf"),
        ("02_ПД", "01_ПЗ", "523-ПИР-24-ПЗ", "volume-1.pdf"),
    ]
    for area, section, designation, filename in files:
        folder = source / area / section / designation / f"{designation}_260728"
        folder.mkdir(parents=True)
        (folder / filename).write_bytes(filename.encode())
    entries = [
        RegEntry(
            oboznachenie="523-ПИР-24-ПЗ", key="ПЗ", area="ПД", razdel="1",
            section="01_ПЗ", tom="1",
            name="Раздел 1. Пояснительная записка"),
        RegEntry(
            oboznachenie="523-ПИР-24-ПЗУ1", key="ПЗУ1", area="ПД", razdel="2",
            section="02_СПОЗУ", tom="2.1",
            name="Раздел 2. Схема планировочной организации земельного участка"),
        RegEntry(
            oboznachenie="523-ПИР-24-ИГИ", key="ИГИ", area="ИИ",
            section="02_ИГИ", tom="2",
            name="Инженерно-геологические изыскания"),
    ]
    target = tmp_path / "package"
    result = customer_package.build_package(
        AppConfig(project_root=str(tmp_path), latest_dir="!LATEST"),
        str(source), str(target), entries=entries)

    wb = load_workbook(result["registry"])
    ws = wb["Состав"]
    titles = [ws.cell(row, 3).value for row in range(7, ws.max_row + 1)]
    assert titles.index("Проектная документация") < titles.index(
        "Отчетная техническая документация")
    document_rows = [
        row for row in range(7, ws.max_row + 1)
        if ws.cell(row, 1).value is not None]
    assert [ws.cell(row, 1).value for row in document_rows] == ["1", "2.1", "2"]
    for title in (
            "Проектная документация",
            "Раздел 1. Пояснительная записка",
            "Отчетная техническая документация"):
        row = next(
            row for row in range(7, ws.max_row + 1)
            if ws.cell(row, 3).value == title)
        assert ws.cell(row, 3).font.bold is True
        assert ws.cell(row, 3).fill.fgColor.rgb.endswith("D9E2F3")
    wb.close()


def test_customer_package_keeps_smeta_parent_folders(tmp_path):
    source = tmp_path / "!LATEST"
    vor = (source / "02_ПД" / "12_СМ" / "519-ПИР-23-СМ2"
           / "519-ПИР-23-СМ2-ПОС-ВОР1_260101")
    lsr = (source / "02_ПД" / "12_СМ" / "519-ПИР-23-СМ3"
           / "519-ПИР-23-СМ3_01-01-01_ЛСР_260101")
    vor.mkdir(parents=True)
    lsr.mkdir(parents=True)
    (vor / "vor.pdf").write_bytes(b"vor")
    (lsr / "lsr.xlsx").write_bytes(b"lsr")
    target = tmp_path / "package"

    customer_package.build_package(
        AppConfig(project_root=str(tmp_path), latest_dir="!LATEST"),
        str(source), str(target),
        selected_folders=[os.path.join("02_ПД", "12_СМ")])

    assert (target / "02_ПД" / "12_СМ" / "519-ПИР-23-СМ2"
            / "519-ПИР-23-СМ2-ПОС-ВОР1_260101" / "vor.pdf").is_file()
    assert (target / "02_ПД" / "12_СМ" / "519-ПИР-23-СМ3"
            / "519-ПИР-23-СМ3_01-01-01_ЛСР_260101" / "lsr.xlsx").is_file()
